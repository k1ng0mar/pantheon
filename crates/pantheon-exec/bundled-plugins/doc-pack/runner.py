#!/usr/bin/env python3
"""doc-pack runner: create and read .docx / .xlsx / .pptx documents.

Speaks Pantheon's tool-plugin JSON protocol over stdio:
    stdin  -> {"call_id": "...", "tool": "...", "args": {...}}
    stdout -> {"call_id": "...", "result": ...}
           or {"call_id": "...", "error": {"code": "...", "cause": "..."}}

Only the Python standard library is imported unconditionally. The three
document libraries are optional per-tool: the runner checks for them at
startup, prints a clear notice to stderr for anything missing, and any
tool whose library is absent fails with DOC_PACK_MISSING_DEP carrying
the exact pip command to run. Partial installs keep working (e.g. docx
tools work without python-pptx).
"""

import json
import os
import sys

# module name -> pip package name
REQUIRED_LIBS = (
    ("docx", "python-docx"),
    ("openpyxl", "openpyxl"),
    ("pptx", "python-pptx"),
)
INSTALL_HINT = "pip install python-docx openpyxl python-pptx"

# tool -> module it needs
TOOL_LIB = {
    "docx_create": "docx",
    "docx_read_text": "docx",
    "xlsx_create": "openpyxl",
    "xlsx_read": "openpyxl",
    "pptx_create": "pptx",
    "pptx_read_text": "pptx",
}


def missing_dependencies():
    """Return the pip package names of the document libs not importable."""
    missing = []
    for module, package in REQUIRED_LIBS:
        try:
            __import__(module)
        except ImportError:
            missing.append(package)
    return missing


def startup_notice():
    missing = missing_dependencies()
    if missing:
        sys.stderr.write(
            "doc-pack: missing document libraries: %s.\n"
            "doc-pack: install them with: %s\n"
            "doc-pack: tools needing a missing library will fail with "
            "DOC_PACK_MISSING_DEP until it is installed.\n"
            % (", ".join(missing), INSTALL_HINT)
        )
        sys.stderr.flush()


def require_lib(tool):
    """Import and return the module for a tool, or raise DocPackError."""
    module = TOOL_LIB[tool]
    try:
        return __import__(module)
    except ImportError:
        package = dict(REQUIRED_LIBS)[module]
        raise DocPackError(
            "DOC_PACK_MISSING_DEP",
            "tool '%s' needs the '%s' library, which is not installed. "
            "Install it with: %s" % (tool, package, INSTALL_HINT),
        )


class DocPackError(Exception):
    def __init__(self, code, cause):
        super().__init__(cause)
        self.code = code
        self.cause = cause


def resolve_out_path(path):
    """Absolute, parent-creating output path. Relative paths are refused."""
    if not isinstance(path, str) or not path:
        raise DocPackError("DOC_PACK_BAD_ARGS", "path must be a non-empty string")
    path = os.path.expanduser(path)
    if not os.path.isabs(path):
        raise DocPackError(
            "DOC_PACK_BAD_PATH",
            "path must be absolute, got %r" % path,
        )
    parent = os.path.dirname(path)
    if parent:
        os.makedirs(parent, exist_ok=True)
    return path


def resolve_in_path(path):
    if not isinstance(path, str) or not path:
        raise DocPackError("DOC_PACK_BAD_ARGS", "path must be a non-empty string")
    path = os.path.expanduser(path)
    if not os.path.isabs(path):
        raise DocPackError(
            "DOC_PACK_BAD_PATH",
            "path must be absolute, got %r" % path,
        )
    if not os.path.isfile(path):
        raise DocPackError("DOC_PACK_IO", "file not found: %s" % path)
    return path


# ---------------------------------------------------------------- docx ---

def tool_docx_create(args):
    docx = require_lib("docx_create")
    from docx.shared import Inches

    path = resolve_out_path(args.get("path"))
    blocks = args.get("blocks")
    if not isinstance(blocks, list):
        raise DocPackError("DOC_PACK_BAD_ARGS", "blocks must be a list of block objects")

    doc = docx.Document()
    if args.get("title"):
        doc.core_properties.title = str(args["title"])
    if args.get("author"):
        doc.core_properties.author = str(args["author"])

    n_paragraphs = 0
    n_tables = 0
    for i, block in enumerate(blocks):
        if not isinstance(block, dict):
            raise DocPackError("DOC_PACK_BAD_ARGS", "block %d is not an object" % i)
        btype = block.get("type")
        if btype == "heading":
            level = block.get("level", 1)
            if not isinstance(level, int) or not 0 <= level <= 9:
                raise DocPackError("DOC_PACK_BAD_ARGS", "block %d: heading level must be 0-9" % i)
            doc.add_heading(str(block.get("text", "")), level=level)
            n_paragraphs += 1
        elif btype == "paragraph":
            doc.add_paragraph(str(block.get("text", "")), style=block.get("style"))
            n_paragraphs += 1
        elif btype in ("bullet", "number"):
            style = "List Bullet" if btype == "bullet" else "List Number"
            level = block.get("level", 0)
            if not isinstance(level, int) or not 0 <= level <= 8:
                raise DocPackError("DOC_PACK_BAD_ARGS", "block %d: list level must be 0-8" % i)
            p = doc.add_paragraph(str(block.get("text", "")), style=style)
            if level:
                p.paragraph_format.left_indent = Inches(0.25 * level)
            n_paragraphs += 1
        elif btype == "table":
            rows = block.get("rows")
            if not isinstance(rows, list) or not rows or not all(isinstance(r, list) for r in rows):
                raise DocPackError("DOC_PACK_BAD_ARGS", "block %d: table needs a non-empty rows list of lists" % i)
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = block.get("style", "Table Grid")
            header = bool(block.get("header", False))
            for ri, row in enumerate(rows):
                for ci, val in enumerate(row):
                    if ci >= len(table.columns):
                        break
                    cell = table.cell(ri, ci)
                    cell.text = "" if val is None else str(val)
                    if header and ri == 0:
                        for par in cell.paragraphs:
                            for run in par.runs:
                                run.bold = True
            n_tables += 1
        elif btype == "page_break":
            doc.add_page_break()
            n_paragraphs += 1
        else:
            raise DocPackError(
                "DOC_PACK_BAD_ARGS",
                "block %d: unknown type %r (heading|paragraph|bullet|number|table|page_break)" % (i, btype),
            )

    doc.save(path)
    return {"path": path, "paragraphs": n_paragraphs, "tables": n_tables}


def tool_docx_read_text(args):
    docx = require_lib("docx_read_text")
    path = resolve_in_path(args.get("path"))
    try:
        doc = docx.Document(path)
    except Exception as e:
        raise DocPackError("DOC_PACK_IO", "cannot open %s: %s" % (path, e))

    paragraphs = []
    for p in doc.paragraphs:
        text = p.text
        if text.strip():
            paragraphs.append({"style": p.style.name if p.style else "", "text": text})
    tables = []
    for t in doc.tables:
        tables.append([[c.text for c in row.cells] for row in t.rows])
    full = "\n".join(p["text"] for p in paragraphs)
    return {"path": path, "paragraphs": paragraphs, "tables": tables, "text": full}


# ---------------------------------------------------------------- xlsx ---

def _xlsx_apply_header(ws, ncols):
    from openpyxl.styles import Alignment, Font, PatternFill

    font = Font(bold=True)
    fill = PatternFill("solid", fgColor="D9D9D9")
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = font
        cell.fill = fill
        cell.alignment = Alignment(vertical="center")


def _xlsx_autofit(ws, rows):
    widths = []
    for row in rows:
        for i, val in enumerate(row):
            s = "" if val is None else str(val)
            if i >= len(widths):
                widths.append(0)
            widths[i] = max(widths[i], len(s))
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(w + 2, 50)


def tool_xlsx_create(args):
    openpyxl = require_lib("xlsx_create")
    path = resolve_out_path(args.get("path"))
    sheets = args.get("sheets")
    if not isinstance(sheets, list) or not sheets:
        raise DocPackError("DOC_PACK_BAD_ARGS", "sheets must be a non-empty list")

    wb = openpyxl.Workbook()
    # Reuse the default sheet for the first entry, rename it.
    first = True
    names = []
    for i, spec in enumerate(sheets):
        if not isinstance(spec, dict):
            raise DocPackError("DOC_PACK_BAD_ARGS", "sheet %d is not an object" % i)
        name = str(spec.get("name", "Sheet%d" % (i + 1)))[:31]
        rows = spec.get("rows", [])
        if not isinstance(rows, list):
            raise DocPackError("DOC_PACK_BAD_ARGS", "sheet %r: rows must be a list of lists" % name)
        ws = wb.active if first else wb.create_sheet(title=name)
        if first:
            ws.title = name
        first = False
        names.append(ws.title)

        for row in rows:
            if not isinstance(row, list):
                raise DocPackError("DOC_PACK_BAD_ARGS", "sheet %r: each row must be a list" % name)
            out = []
            for val in row:
                if isinstance(val, str) and val.startswith("=") and len(val) > 1:
                    out.append(val)  # openpyxl treats leading "=" as a formula
                else:
                    out.append(val)
            ws.append(out)

        ncols = max((len(r) for r in rows), default=0)
        if spec.get("header") and rows:
            _xlsx_apply_header(ws, ncols)
        widths = spec.get("col_widths")
        if isinstance(widths, list) and widths:
            for ci, w in enumerate(widths, start=1):
                try:
                    ws.column_dimensions[ws.cell(row=1, column=ci).column_letter].width = float(w)
                except (TypeError, ValueError):
                    pass
        else:
            _xlsx_autofit(ws, rows)
        freeze = spec.get("freeze")
        if isinstance(freeze, str) and freeze:
            ws.freeze_panes = freeze

    wb.save(path)
    return {"path": path, "sheets": names}


def tool_xlsx_read(args):
    openpyxl = require_lib("xlsx_read")
    path = resolve_in_path(args.get("path"))
    try:
        wb = openpyxl.load_workbook(path, data_only=False, read_only=True)
    except Exception as e:
        raise DocPackError("DOC_PACK_IO", "cannot open %s: %s" % (path, e))
    names = wb.sheetnames
    wanted = args.get("sheet")
    if wanted and wanted not in names:
        raise DocPackError("DOC_PACK_BAD_ARGS", "no sheet named %r (have: %s)" % (wanted, ", ".join(names)))
    ws = wb[wanted] if wanted else wb.active
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append(["" if v is None else v for v in row])
    # Drop fully-empty trailing rows for a cleaner result.
    while rows and all(v == "" for v in rows[-1]):
        rows.pop()
    return {"path": path, "sheets": names, "sheet": ws.title, "rows": rows}


# ---------------------------------------------------------------- pptx ---

# Default-template layout indices (python-pptx ships these).
PPTX_LAYOUTS = {
    "title_slide": 0,
    "title_content": 1,
    "section_header": 2,
    "two_content": 3,
    "comparison": 4,
    "title_only": 5,
    "blank": 6,
}


def tool_pptx_create(args):
    pptx = require_lib("pptx_create")
    path = resolve_out_path(args.get("path"))
    slides = args.get("slides")
    if not isinstance(slides, list) or not slides:
        raise DocPackError("DOC_PACK_BAD_ARGS", "slides must be a non-empty list")

    prs = pptx.Presentation()
    if args.get("title"):
        prs.core_properties.title = str(args["title"])

    for i, spec in enumerate(slides):
        if not isinstance(spec, dict):
            raise DocPackError("DOC_PACK_BAD_ARGS", "slide %d is not an object" % i)
        layout_name = spec.get("layout", "title_content")
        layout_idx = PPTX_LAYOUTS.get(layout_name)
        if layout_idx is None:
            raise DocPackError(
                "DOC_PACK_BAD_ARGS",
                "slide %d: unknown layout %r (one of: %s)"
                % (i, layout_name, ", ".join(sorted(PPTX_LAYOUTS))),
            )
        if layout_idx >= len(prs.slide_layouts):
            raise DocPackError("DOC_PACK_BAD_ARGS", "slide %d: layout %r not in template" % (i, layout_name))
        slide = prs.slides.add_slide(prs.slide_layouts[layout_idx])

        title = spec.get("title")
        if title and slide.shapes.title is not None:
            slide.shapes.title.text = str(title)

        bullets = spec.get("bullets", [])
        if bullets:
            if not isinstance(bullets, list):
                raise DocPackError("DOC_PACK_BAD_ARGS", "slide %d: bullets must be a list" % i)
            body = None
            for ph in slide.placeholders:
                if ph.has_text_frame and ph != slide.shapes.title:
                    body = ph
                    break
            if body is None:
                raise DocPackError("DOC_PACK_BAD_ARGS", "slide %d: layout %r has no content placeholder" % (i, layout_name))
            tf = body.text_frame
            tf.clear()
            for bi, item in enumerate(bullets):
                if isinstance(item, dict):
                    text, level = str(item.get("text", "")), int(item.get("level", 0))
                else:
                    text, level = str(item), 0
                if not 0 <= level <= 8:
                    raise DocPackError("DOC_PACK_BAD_ARGS", "slide %d bullet %d: level must be 0-8" % (i, bi))
                p = tf.paragraphs[0] if bi == 0 else tf.add_paragraph()
                p.text = text
                p.level = level

        notes = spec.get("notes")
        if notes:
            slide.notes_slide.placeholders[1].text = str(notes)

    prs.save(path)
    return {"path": path, "slides": len(slides)}


def tool_pptx_read_text(args):
    pptx = require_lib("pptx_read_text")
    path = resolve_in_path(args.get("path"))
    try:
        prs = pptx.Presentation(path)
    except Exception as e:
        raise DocPackError("DOC_PACK_IO", "cannot open %s: %s" % (path, e))

    slides = []
    for idx, slide in enumerate(prs.slides):
        title_shape = slide.shapes.title
        texts = []
        for shape in slide.shapes:
            if not shape.has_text_frame or shape == title_shape:
                continue
            t = shape.text.strip()
            if t:
                texts.append(t)
        notes = ""
        if slide.has_notes_slide:
            for shape in slide.notes_slide.shapes:
                if shape.has_text_frame and shape.text.strip():
                    notes = shape.text.strip()
                    break
        slides.append({
            "index": idx,
            "title": title_shape.text if title_shape is not None else "",
            "texts": texts,
            "notes": notes,
        })
    return {"path": path, "slides": slides}


TOOLS = {
    "docx_create": tool_docx_create,
    "docx_read_text": tool_docx_read_text,
    "xlsx_create": tool_xlsx_create,
    "xlsx_read": tool_xlsx_read,
    "pptx_create": tool_pptx_create,
    "pptx_read_text": tool_pptx_read_text,
}


def handle_request(req):
    call_id = req.get("call_id", "unknown")
    tool = req.get("tool")
    args = req.get("args") or {}
    if not isinstance(args, dict):
        return {"call_id": call_id,
                "error": {"code": "DOC_PACK_BAD_ARGS", "cause": "args must be an object"}}
    fn = TOOLS.get(tool)
    if fn is None:
        return {"call_id": call_id,
                "error": {"code": "DOC_PACK_UNKNOWN_TOOL",
                          "cause": "unknown tool %r (have: %s)" % (tool, ", ".join(sorted(TOOLS)))}}
    try:
        return {"call_id": call_id, "result": fn(args)}
    except DocPackError as e:
        return {"call_id": call_id, "error": {"code": e.code, "cause": e.cause}}
    except Exception as e:
        return {"call_id": call_id,
                "error": {"code": "DOC_PACK_INTERNAL",
                          "cause": "%s: %s" % (type(e).__name__, e)}}


def main():
    startup_notice()
    out = sys.stdout
    for line in sys.stdin:
        line = line.strip()
        if not line:
            continue
        try:
            req = json.loads(line)
        except json.JSONDecodeError as e:
            out.write(json.dumps({"call_id": "unknown",
                                  "error": {"code": "DOC_PACK_BAD_ARGS",
                                            "cause": "request is not valid JSON: %s" % e}}) + "\n")
            out.flush()
            continue
        if not isinstance(req, dict):
            out.write(json.dumps({"call_id": "unknown",
                                  "error": {"code": "DOC_PACK_BAD_ARGS",
                                            "cause": "request must be a JSON object"}}) + "\n")
            out.flush()
            continue
        out.write(json.dumps(handle_request(req), default=str) + "\n")
        out.flush()


if __name__ == "__main__":
    main()
